from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from surveying.domain.error_analysis import ObservationStats, assess_closure
from surveying.domain.leveling import LevelingResult
from surveying.domain.models import DomainError, LevelingProject


class ExportError(DomainError):
    pass


def _mm(value_m: float, digits: int = 2) -> str:
    return f"{value_m * 1000.0:.{digits}f}"


def _fixed(value: float, digits: int = 4) -> str:
    return f"{value:.{digits}f}"


def _conclusion(result: LevelingResult) -> str:
    return "合格" if result.is_within_tolerance else "超限"


def _quality_label(result: LevelingResult) -> str:
    return result.quality_level.value


def _observation_rows(stats: ObservationStats) -> list[list[object]]:
    reference = stats.target if stats.target is not None else stats.mean
    rows: list[list[object]] = []
    outlier_indexes = {item.index for item in stats.outliers}
    for index, value in enumerate(stats.values, start=1):
        rows.append(
            [
                index,
                _fixed(value),
                _fixed(value - reference, 5),
                "粗差可疑" if index in outlier_indexes else "正常",
            ]
        )
    return rows


class ResultExporter:
    @staticmethod
    def export_csv(
        path: str | Path,
        project: LevelingProject,
        result: LevelingResult,
        stats: ObservationStats | None = None,
    ) -> None:
        destination = Path(path)
        closure = assess_closure(result)
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            with destination.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["测量数据处理成果"])
                writer.writerow(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
                writer.writerow([])
                writer.writerow(["一、工程概况"])
                writer.writerow(["工程名称", project.name])
                writer.writerow(["路线类型", project.route_type.label])
                writer.writerow(["测量等级", project.grade.value])
                writer.writerow(["观测者", project.observer])
                writer.writerow(["观测日期", project.survey_date])
                writer.writerow(["起点", project.start_point, "起点高程(m)", _fixed(project.start_elevation_m)])
                writer.writerow(["终点", result.end_point, "终点高程(m)", _fixed(result.target_end_elevation_m)])
                writer.writerow(["备注", project.notes])
                writer.writerow([])
                writer.writerow(["二、计算摘要"])
                writer.writerow(["路线总长(m)", _fixed(result.total_distance_m, 3)])
                writer.writerow(["观测高差之和(m)", _fixed(result.observed_difference_sum_m)])
                writer.writerow(["改正数之和(mm)", _mm(result.correction_sum_m)])
                writer.writerow(["闭合差(mm)", _mm(result.closure_error_m)])
                writer.writerow(["容许闭合差(mm)", _mm(result.allowable_closure_m)])
                writer.writerow(
                    [
                        "闭合差利用率",
                        "—" if closure.utilization is None else f"{closure.utilization * 100:.1f}%",
                    ]
                )
                writer.writerow(["闭合差与路线长之比", closure.ratio_text])
                writer.writerow(["质量评定", _quality_label(result)])
                writer.writerow(["计算结论", _conclusion(result)])
                writer.writerow(["改正数中误差(mm)", f"{closure.correction_rms_mm:.3f}"])
                writer.writerow(["最大改正数(mm)", f"{closure.max_abs_correction_mm:.3f}"])
                writer.writerow(["计算终点高程(m)", _fixed(result.computed_end_elevation_m)])
                writer.writerow([])
                writer.writerow(["三、测段计算成果"])
                writer.writerow(
                    [
                        "序号",
                        "起点",
                        "终点",
                        "距离(m)",
                        "观测高差(m)",
                        "改正数(mm)",
                        "改正后高差(m)",
                        "起点高程(m)",
                        "终点高程(m)",
                        "备注",
                    ]
                )
                for item in result.segments:
                    writer.writerow(
                        [
                            item.index,
                            item.segment.from_point,
                            item.segment.to_point,
                            _fixed(item.segment.distance_m, 3),
                            _fixed(item.segment.observed_difference_m),
                            f"{item.correction_mm:.3f}",
                            _fixed(item.corrected_difference_m),
                            _fixed(item.from_elevation_m),
                            _fixed(item.to_elevation_m),
                            item.segment.remark,
                        ]
                    )
                writer.writerow([])
                writer.writerow(["四、误差分析"])
                if stats is None:
                    writer.writerow(["观测列统计", "未进行重复观测分析"])
                else:
                    writer.writerow(["观测个数", stats.count])
                    writer.writerow(["观测均值", _fixed(stats.mean, 5)])
                    writer.writerow(["中位数", _fixed(stats.median, 5)])
                    writer.writerow(["样本标准差", _fixed(stats.std_dev, 5)])
                    writer.writerow(["标准误差", _fixed(stats.standard_error, 5)])
                    writer.writerow(["最小值", _fixed(stats.minimum, 5)])
                    writer.writerow(["最大值", _fixed(stats.maximum, 5)])
                    writer.writerow(["极差", _fixed(stats.range_value, 5)])
                    writer.writerow(["RMSE", _fixed(stats.rmse, 5)])
                    writer.writerow(
                        [
                            "相对标准差",
                            "—"
                            if stats.relative_std is None
                            else f"{stats.relative_std * 100:.3f}%",
                        ]
                    )
                    writer.writerow(
                        [
                            "参考值",
                            "—" if stats.target is None else _fixed(stats.target, 5),
                        ]
                    )
                    writer.writerow(
                        [
                            "粗差判定",
                            f"|残差| > {stats.outlier_sigma:g}σ，"
                            f"共 {len(stats.outliers)} 个可疑值",
                        ]
                    )
                    writer.writerow([])
                    writer.writerow(["序号", "观测值", "残差", "判定"])
                    for row in _observation_rows(stats):
                        writer.writerow(row)
                if result.warnings:
                    writer.writerow([])
                    writer.writerow(["五、提示信息"])
                    for warning in result.warnings:
                        writer.writerow([warning])
        except OSError as exc:
            raise ExportError(f"导出 CSV 失败：{exc}") from exc

    @staticmethod
    def export_excel(
        path: str | Path,
        project: LevelingProject,
        result: LevelingResult,
        stats: ObservationStats | None = None,
    ) -> None:
        destination = Path(path)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ExportError("缺少 openpyxl，无法导出 Excel。请先安装项目依赖。") from exc

        navy = "17324D"
        teal = "0F766E"
        pale = "E8F2F1"
        light = "F4F7F9"
        white = "FFFFFF"
        border_color = "D8E1E8"
        thin = Side(style="thin", color=border_color)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        title_font = Font(name="Microsoft YaHei", size=16, bold=True, color=navy)
        section_font = Font(name="Microsoft YaHei", size=11, bold=True, color=white)
        header_font = Font(name="Microsoft YaHei", size=10, bold=True, color=white)
        body_font = Font(name="Microsoft YaHei", size=10, color="263238")
        subtle_font = Font(name="Microsoft YaHei", size=9, color="607D8B")
        section_fill = PatternFill("solid", fgColor=teal)
        header_fill = PatternFill("solid", fgColor=navy)
        pale_fill = PatternFill("solid", fgColor=pale)
        light_fill = PatternFill("solid", fgColor=light)
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        def section(sheet, row: int, title: str, width: int) -> int:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            cell = sheet.cell(row=row, column=1, value=title)
            cell.font = section_font
            cell.fill = section_fill
            cell.alignment = left
            sheet.row_dimensions[row].height = 24
            return row + 1

        def table_header(sheet, row: int, headers: list[str]) -> int:
            for column, header in enumerate(headers, start=1):
                cell = sheet.cell(row=row, column=column, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = border
            sheet.row_dimensions[row].height = 26
            return row + 1

        def body_row(sheet, row: int, values: list[object], number_formats: dict[int, str]) -> int:
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = body_font
                cell.border = border
                cell.alignment = left if column in (1, 2, 3) else center
                if column in number_formats:
                    cell.number_format = number_formats[column]
                if row % 2 == 0:
                    cell.fill = light_fill
            return row + 1

        workbook = Workbook()
        summary = workbook.active
        summary.title = "工程概况"
        summary.sheet_view.showGridLines = False
        summary.merge_cells("A1:D1")
        title_cell = summary["A1"]
        title_cell.value = project.name or "测量数据处理成果"
        title_cell.font = title_font
        title_cell.fill = pale_fill
        title_cell.alignment = center
        summary.row_dimensions[1].height = 38
        summary.merge_cells("A2:D2")
        subtitle = summary["A2"]
        subtitle.value = (
            f"{project.route_type.label}｜{project.grade.value}｜"
            f"导出时间 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        subtitle.font = subtle_font
        subtitle.alignment = center

        row = 4
        row = section(summary, row, "工程信息", 4)
        metadata = [
            ("工程名称", project.name, "路线类型", project.route_type.label),
            ("测量等级", project.grade.value, "观测者", project.observer or "—"),
            ("观测日期", project.survey_date, "备注", project.notes or "—"),
            ("起点名称", project.start_point, "起点高程(m)", float(project.start_elevation_m)),
            ("终点名称", result.end_point, "终点高程(m)", float(result.target_end_elevation_m)),
        ]
        for item in metadata:
            for column, value in enumerate(item, start=1):
                cell = summary.cell(row=row, column=column, value=value)
                cell.font = body_font
                cell.border = border
                if column in (1, 3):
                    cell.fill = pale_fill
                    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=navy)
                if column == 4 and isinstance(value, float):
                    cell.number_format = "0.0000"
            summary.row_dimensions[row].height = 24
            row += 1

        row += 1
        row = section(summary, row, "计算摘要", 4)
        closure = assess_closure(result)
        summary_rows = [
            ("路线总长(m)", float(result.total_distance_m), "观测高差之和(m)", float(result.observed_difference_sum_m)),
            ("闭合差(mm)", float(result.closure_error_mm), "容许闭合差(mm)", float(result.allowable_closure_mm)),
            (
                "闭合差利用率",
                "—" if closure.utilization is None else f"{closure.utilization * 100:.1f}%",
                "闭合差与路线长之比",
                closure.ratio_text,
            ),
            ("质量评定", result.quality_level.value, "计算结论", _conclusion(result)),
            ("改正数中误差(mm)", float(closure.correction_rms_mm), "最大改正数(mm)", float(closure.max_abs_correction_mm)),
            ("计算终点高程(m)", float(result.computed_end_elevation_m), "测段数量", len(result.segments)),
        ]
        for item in summary_rows:
            for column, value in enumerate(item, start=1):
                cell = summary.cell(row=row, column=column, value=value)
                cell.font = body_font
                cell.border = border
                if column in (1, 3):
                    cell.fill = pale_fill
                    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=navy)
                if column in (2, 4) and isinstance(value, float):
                    cell.number_format = "0.0000"
            summary.row_dimensions[row].height = 24
            row += 1

        if result.warnings:
            row += 1
            row = section(summary, row, "提示信息", 4)
            for warning in result.warnings:
                summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                cell = summary.cell(row=row, column=1, value=warning)
                cell.font = Font(name="Microsoft YaHei", size=9, color="B45309")
                cell.alignment = left
                cell.fill = PatternFill("solid", fgColor="FFF7E6")
                cell.border = border
                row += 1

        summary.column_dimensions["A"].width = 20
        summary.column_dimensions["B"].width = 28
        summary.column_dimensions["C"].width = 24
        summary.column_dimensions["D"].width = 32

        calc = workbook.create_sheet("水准计算")
        calc.sheet_view.showGridLines = False
        calc.merge_cells("A1:J1")
        calc["A1"].value = "测段计算成果"
        calc["A1"].font = title_font
        calc["A1"].fill = pale_fill
        calc["A1"].alignment = center
        calc.row_dimensions[1].height = 34
        headers = [
            "序号",
            "起点",
            "终点",
            "距离(m)",
            "观测高差(m)",
            "改正数(mm)",
            "改正后高差(m)",
            "起点高程(m)",
            "终点高程(m)",
            "备注",
        ]
        row = table_header(calc, 3, headers)
        number_formats = {
            4: "0.000",
            5: "0.0000",
            6: "0.000",
            7: "0.0000",
            8: "0.0000",
            9: "0.0000",
        }
        for item in result.segments:
            values: list[object] = [
                item.index,
                item.segment.from_point,
                item.segment.to_point,
                float(item.segment.distance_m),
                float(item.segment.observed_difference_m),
                float(item.correction_mm),
                float(item.corrected_difference_m),
                float(item.from_elevation_m),
                float(item.to_elevation_m),
                item.segment.remark or "",
            ]
            row = body_row(calc, row, values, number_formats)
        widths = [8, 14, 14, 13, 15, 14, 17, 16, 16, 28]
        for index, width in enumerate(widths, start=1):
            calc.column_dimensions[get_column_letter(index)].width = width
        calc.freeze_panes = "A4"

        analysis = workbook.create_sheet("误差分析")
        analysis.sheet_view.showGridLines = False
        analysis.merge_cells("A1:F1")
        analysis["A1"].value = "误差分析"
        analysis["A1"].font = title_font
        analysis["A1"].fill = pale_fill
        analysis["A1"].alignment = center
        analysis.row_dimensions[1].height = 34
        row = 3
        row = section(analysis, row, "闭合差评定", 6)
        closure_rows = [
            ("闭合差(mm)", float(result.closure_error_mm), "容许闭合差(mm)", float(result.allowable_closure_mm)),
            (
                "闭合差利用率",
                "—" if closure.utilization is None else f"{closure.utilization * 100:.1f}%",
                "结论",
                _conclusion(result),
            ),
            ("质量评定", result.quality_level.value, "闭合差/路线长", result.misclosure_ratio_text),
            ("改正数中误差(mm)", float(closure.correction_rms_mm), "最大改正数(mm)", float(closure.max_abs_correction_mm)),
        ]
        for item in closure_rows:
            for column, value in enumerate(item, start=1):
                cell = analysis.cell(row=row, column=column, value=value)
                cell.font = body_font
                cell.border = border
                if column in (1, 3):
                    cell.fill = pale_fill
                    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=navy)
                if column in (2, 4) and isinstance(value, float):
                    cell.number_format = "0.000"
            analysis.row_dimensions[row].height = 24
            row += 1

        row += 1
        row = section(analysis, row, "重复观测统计", 6)
        if stats is None:
            analysis.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell = analysis.cell(row=row, column=1, value="未进行重复观测分析")
            cell.font = subtle_font
            cell.alignment = left
            row += 1
        else:
            stat_rows = [
                ("观测个数", stats.count, "参考值", "—" if stats.target is None else float(stats.target)),
                ("观测均值", float(stats.mean), "中位数", float(stats.median)),
                ("样本标准差", float(stats.std_dev), "标准误差", float(stats.standard_error)),
                ("最小值", float(stats.minimum), "最大值", float(stats.maximum)),
                ("极差", float(stats.range_value), "RMSE", float(stats.rmse)),
                (
                    "相对标准差",
                    "—" if stats.relative_std is None else f"{stats.relative_std * 100:.3f}%",
                    "粗差判定倍数",
                    f"{stats.outlier_sigma:g}σ",
                ),
                ("可疑值数量", len(stats.outliers), "判定方法", "残差绝对值大于判定倍数×标准差"),
            ]
            for item in stat_rows:
                for column, value in enumerate(item, start=1):
                    cell = analysis.cell(row=row, column=column, value=value)
                    cell.font = body_font
                    cell.border = border
                    if column in (1, 3):
                        cell.fill = pale_fill
                        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=navy)
                    if isinstance(value, float):
                        cell.number_format = "0.00000"
                analysis.row_dimensions[row].height = 24
                row += 1

            row += 1
            row = section(analysis, row, "观测值与残差", 6)
            row = table_header(analysis, row, ["序号", "观测值", "残差", "判定", "", ""])
            outlier_indexes = {item.index for item in stats.outliers}
            reference = stats.target if stats.target is not None else stats.mean
            for index, value in enumerate(stats.values, start=1):
                values = [
                    index,
                    float(value),
                    float(value - reference),
                    "粗差可疑" if index in outlier_indexes else "正常",
                    "",
                    "",
                ]
                for column, cell_value in enumerate(values, start=1):
                    cell = analysis.cell(row=row, column=column, value=cell_value)
                    cell.font = body_font
                    cell.border = border
                    cell.alignment = center
                    if column in (2, 3):
                        cell.number_format = "0.00000"
                    if index in outlier_indexes:
                        cell.fill = PatternFill("solid", fgColor="FDE8E8")
                row += 1

        widths = [22, 18, 22, 22, 12, 12]
        for index, width in enumerate(widths, start=1):
            analysis.column_dimensions[get_column_letter(index)].width = width
        analysis.freeze_panes = "A3"

        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(destination)
        except OSError as exc:
            raise ExportError(f"导出 Excel 失败：{exc}") from exc
