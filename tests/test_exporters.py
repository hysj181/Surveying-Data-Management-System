from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from surveying.domain.error_analysis import analyze_observations
from surveying.domain.leveling import calculate_leveling
from surveying.services.exporters import ResultExporter
from surveying.services.samples import create_sample_project


def test_csv_export(tmp_path: Path) -> None:
    project = create_sample_project()
    result = calculate_leveling(project)
    stats = analyze_observations([1.5234, 1.5238, 1.5240, 1.5241], target=1.5238)
    path = tmp_path / "result.csv"

    ResultExporter.export_csv(path, project, result, stats)

    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert "滨河路四等水准测量" in text
    assert "-8.00" in text
    assert "测段计算成果" in text
    assert "样本标准差" in text
    assert "BM-A" in text


def test_excel_export(tmp_path: Path) -> None:
    project = create_sample_project()
    result = calculate_leveling(project)
    stats = analyze_observations([1.5234, 1.5238, 1.5240, 1.5241], target=1.5238)
    path = tmp_path / "result.xlsx"

    ResultExporter.export_excel(path, project, result, stats)

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["工程概况", "水准计算", "误差分析"]
    assert workbook["工程概况"]["A1"].value == "滨河路四等水准测量"
    assert workbook["水准计算"].max_row == 8
    assert workbook["误差分析"]["A1"].value == "误差分析"


def test_export_without_observations(tmp_path: Path) -> None:
    project = create_sample_project()
    result = calculate_leveling(project)
    csv_path = tmp_path / "no-obs.csv"
    xlsx_path = tmp_path / "no-obs.xlsx"

    ResultExporter.export_csv(csv_path, project, result)
    ResultExporter.export_excel(xlsx_path, project, result)

    assert "未进行重复观测分析" in csv_path.read_text(encoding="utf-8-sig")
    workbook = load_workbook(xlsx_path)
    assert workbook["误差分析"]["A1"].value == "误差分析"
